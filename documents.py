import uuid
from io import BytesIO

from docx import Document as DocxDocument
from llama_index.core import Document, StorageContext, VectorStoreIndex
from llama_index.core.node_parser import SentenceSplitter
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}
CHUNK_SIZE = 512
CHUNK_OVERLAP = 50


def extract_text(filename: str, content: bytes) -> str:
    name = filename.lower()
    if name.endswith((".txt", ".md")):
        return content.decode("utf-8", errors="ignore")
    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if name.endswith(".docx"):
        doc = DocxDocument(BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"მხარდაუჭერელი ფორმატი: {filename}")


def ingest_document(vector_store, filename: str, content: bytes) -> dict:
    text = extract_text(filename, content).strip()
    if not text:
        raise ValueError("ფაილიდან ტექსტი ვერ ამოიღო")

    document_id = str(uuid.uuid4())
    doc = Document(
        text=text,
        metadata={"document_id": document_id, "title": filename, "source": filename},
    )
    splitter = SentenceSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    nodes = splitter.get_nodes_from_documents([doc])

    storage = StorageContext.from_defaults(vector_store=vector_store)
    VectorStoreIndex(nodes=nodes, storage_context=storage)
    return {"document_id": document_id, "filename": filename, "chunks": len(nodes)}


def delete_document(client, collection_name: str, document_id: str) -> None:
    from qdrant_client.http import models

    client.delete(
        collection_name=collection_name,
        points_selector=models.FilterSelector(
            filter=models.Filter(
                must=[models.FieldCondition(key="document_id", match=models.MatchValue(value=document_id))]
            )
        ),
    )
